"""Generate the synthetic test corpus: 30 fictitious Spanish business documents
(10 invoices, 10 contracts, 10 resumes) across PDF, DOCX and TXT, plus
ground_truth.json (expected category + fields) and rag_eval.json (Q&A pairs).

Fully reproducible: seeded Faker (es_CO) + seeded random. All names, companies,
IDs and figures are FICTITIOUS — no real personal data.

Run:  python scripts/generate_corpus.py
"""
import io
import json
import random
import sys
from pathlib import Path

from docx import Document as DocxDocument
from faker import Faker
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

SEED = 42
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "sample_documents"

fake = Faker("es_CO")
Faker.seed(SEED)
rng = random.Random(SEED)

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
         "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def fmt_cop(value: float) -> str:
    return f"{value:,.0f}".replace(",", ".")


def fmt_usd(value: float) -> str:
    return f"{value:,.2f}"


def date_slash(y, m, d) -> str:
    return f"{d:02d}/{m:02d}/{y}"


def date_larga(y, m, d) -> str:
    return f"{d} de {MESES[m - 1]} de {y}"


# ============================================================ invoices
SERVICIOS = [
    "Licencias de software empresarial", "Servicio de consultoria tecnologica",
    "Mantenimiento de equipos de computo", "Servicio de hosting y dominio anual",
    "Desarrollo de modulo de inventarios", "Capacitacion en seguridad informatica",
    "Suministro de papeleria corporativa", "Servicio de aseo y cafeteria",
    "Publicidad digital y pauta en redes", "Soporte tecnico especializado",
    "Arrendamiento de equipos audiovisuales", "Adecuacion de redes electricas",
]


def build_invoice(i: int) -> dict:
    proveedor = fake.company()
    numero = f"FV-{2024 + (i % 2)}-{100 + i * 7}"
    y = 2024 + (i % 2)
    m = rng.randint(1, 12)
    d = rng.randint(1, 28)
    vend_d = min(d + rng.choice([15, 30, 45]), 28)
    vend_m = m + 1 if m < 12 else 1
    vend_y = y if m < 12 else y + 1
    moneda = rng.choice(["COP", "COP", "COP", "COP", "COP", "COP", "USD", "USD", "EUR", "COP"])
    n_items = rng.randint(1, 5)
    items = []
    subtotal = 0.0
    for _ in range(n_items):
        servicio = rng.choice(SERVICIOS)
        cantidad = rng.randint(1, 10)
        if moneda == "COP":
            unitario = rng.randint(80, 4000) * 1000
        else:
            unitario = round(rng.uniform(50, 2500), 2)
        total_item = round(cantidad * unitario, 2)
        subtotal += total_item
        items.append((servicio, cantidad, unitario, total_item))
    subtotal = round(subtotal, 2)
    iva = round(subtotal * 0.19, 2)
    total = round(subtotal + iva, 2)
    fmt = fmt_cop if moneda == "COP" else fmt_usd

    # Edge cases: some invoices omit due date (i==7) or use long date format (i==4)
    fecha_emision = date_larga(y, m, d) if i == 4 else date_slash(y, m, d)
    incluir_vencimiento = i != 7
    fecha_vencimiento = date_slash(vend_y, vend_m, vend_d) if incluir_vencimiento else None

    lines = [
        f"{proveedor}",
        f"NIT: {rng.randint(800, 901)}.{rng.randint(100, 999)}.{rng.randint(100, 999)}-{rng.randint(0, 9)}",
        "Regimen comun - Responsable de IVA",
        "",
        f"FACTURA DE VENTA No: {numero}",
        f"Fecha de emision: {fecha_emision}",
    ]
    if incluir_vencimiento:
        lines.append(f"Fecha de vencimiento: {fecha_vencimiento}")
    cliente = fake.company()
    lines += [
        f"Cliente: {cliente}",
        f"Direccion: {fake.street_address()}, {fake.city()}",
        "",
        "DETALLE DE LA FACTURA",
    ]
    for servicio, cantidad, unitario, total_item in items:
        lines.append(f"- {servicio} | Cantidad: {cantidad} | Valor unitario: ${fmt(unitario)} | Valor: ${fmt(total_item)}")
    lines += [
        "",
        f"Subtotal: ${fmt(subtotal)}",
        f"IVA (19%): ${fmt(iva)}",
        f"Total a pagar: ${fmt(total)}",
        f"Moneda: {moneda}",
        "",
        "Forma de pago: " + rng.choice(["Transferencia bancaria", "Credito 30 dias", "Contado"]),
        "Esta factura se asimila en sus efectos a una letra de cambio (Art. 774 C. de Co.).",
        "Documento generado con datos ficticios para fines academicos.",
    ]
    return {
        "text": "\n".join(lines),
        "title": f"Factura {numero}",
        "ground_truth": {
            "category": "financiero",
            "fields": {
                "proveedor": proveedor,
                "numero_factura": numero,
                "fecha_emision": fecha_emision,
                "fecha_vencimiento": fecha_vencimiento,
                "subtotal": subtotal,
                "impuestos": iva,
                "total": total,
                "moneda": moneda,
            },
        },
    }


# ============================================================ contracts
OBJETOS = [
    "la prestacion de servicios profesionales de desarrollo de software a la medida",
    "el arrendamiento de un local comercial ubicado en el centro de la ciudad",
    "la prestacion de servicios de consultoria en transformacion digital",
    "el suministro periodico de insumos de oficina y papeleria",
    "la prestacion de servicios de soporte y mantenimiento de infraestructura tecnologica",
    "la ejecucion del proyecto de implementacion del sistema de gestion documental",
    "la prestacion de servicios de capacitacion empresarial en habilidades digitales",
    "el diseno, desarrollo e implementacion de la tienda virtual corporativa",
    "la prestacion de servicios de vigilancia y seguridad privada",
    "el licenciamiento y soporte de la plataforma de nomina electronica",
]

OBLIGACIONES_POOL = [
    "Ejecutar el objeto contractual con la debida diligencia y calidad profesional.",
    "Guardar estricta confidencialidad sobre la informacion conocida en desarrollo del contrato.",
    "Entregar informes mensuales de avance a la supervision del contrato.",
    "Cumplir con las normas de seguridad y salud en el trabajo aplicables.",
    "Pagar oportunamente el valor pactado en la forma acordada.",
    "Suministrar la informacion necesaria para la correcta ejecucion del contrato.",
    "Mantener vigentes las polizas y garantias exigidas durante todo el plazo.",
    "Designar un supervisor responsable del seguimiento del contrato.",
    "Devolver los bienes recibidos en el estado en que fueron entregados.",
    "Abstenerse de ceder el contrato sin autorizacion previa y escrita.",
]


def build_contract(i: int) -> dict:
    contratante = fake.company()
    contratista = fake.company() if i % 3 != 0 else f"{fake.first_name()} {fake.last_name()} {fake.last_name()}"
    objeto = OBJETOS[i]
    y = 2024 + (i % 2)
    m = rng.randint(1, 10)
    d = rng.randint(1, 28)
    duracion_meses = rng.choice([3, 6, 12, 12, 24, 36])
    end_m = (m + duracion_meses - 1) % 12 + 1
    end_y = y + (m + duracion_meses - 1) // 12
    fecha_inicio = date_slash(y, m, d)
    fecha_fin = date_slash(end_y, end_m, d)
    if i == 6:
        valor = None  # edge case: contract without explicit value
    else:
        monto = rng.randint(5, 500) * 1_000_000
        valor = f"${fmt_cop(monto)} COP"
    duracion = f"{duracion_meses} meses"
    n_oblig = rng.randint(3, 5)
    obligaciones = rng.sample(OBLIGACIONES_POOL, n_oblig)
    numero = f"CT-{y}-{40 + i * 3}"

    lines = [
        f"CONTRATO No {numero}",
        rng.choice(["CONTRATO DE PRESTACION DE SERVICIOS", "CONTRATO DE ARRENDAMIENTO",
                    "CONTRATO DE SUMINISTRO"]) if i != 1 else "CONTRATO DE ARRENDAMIENTO",
        "",
        f"Contratante: {contratante}",
        f"Contratista: {contratista}",
        "",
        f"Entre los suscritos, {contratante}, sociedad legalmente constituida, quien en adelante se "
        f"denominara EL CONTRATANTE, y {contratista}, quien en adelante se denominara EL CONTRATISTA, "
        "se ha celebrado el presente contrato que se regira por las siguientes clausulas:",
        "",
        f"PRIMERA - OBJETO DEL CONTRATO: {objeto[0].upper() + objeto[1:]}.",
        f"SEGUNDA - PLAZO: El presente contrato tendra una duracion: {duracion}.",
        f"Fecha de inicio: {fecha_inicio}",
        f"Fecha de terminacion: {fecha_fin}",
    ]
    if valor:
        lines.append(f"TERCERA - VALOR: El valor total del contrato es de Valor del contrato: {valor}.")
    lines += [
        "CUARTA - OBLIGACIONES DE LAS PARTES:",
    ]
    for k, ob in enumerate(obligaciones, start=1):
        lines.append(f"{k}. {ob}")
    lines += [
        "QUINTA - TERMINACION: El contrato podra darse por terminado por mutuo acuerdo, por "
        "incumplimiento de las obligaciones o por vencimiento del plazo pactado.",
        "SEXTA - CONFIDENCIALIDAD: Las partes se obligan a mantener reserva sobre la informacion "
        "intercambiada con ocasion de este contrato.",
        "",
        f"Para constancia se firma en {fake.city()} el {date_larga(y, m, d)}.",
        "",
        "Documento generado con datos ficticios para fines academicos.",
    ]
    return {
        "text": "\n".join(lines),
        "title": f"Contrato {numero}",
        "ground_truth": {
            "category": "legal",
            "fields": {
                "partes": [contratante, contratista],
                "objeto_contiene": objeto.split()[3],  # keyword the extraction must contain
                "fecha_inicio": fecha_inicio,
                "fecha_terminacion": fecha_fin,
                "duracion": duracion,
                "valor": valor,
                "num_obligaciones_min": 3,
            },
        },
    }


# ============================================================ resumes
PERFILES = [
    ("Desarrollador de Software", ["Python", "JavaScript", "SQL", "Git", "Docker"],
     "Tecnologo en Desarrollo de Software con experiencia en aplicaciones web empresariales."),
    ("Analista de Datos", ["Power BI", "SQL", "Excel avanzado", "Python", "ETL"],
     "Profesional en analisis de datos orientado a la generacion de reportes e indicadores."),
    ("Ingeniera de Sistemas", ["Java", "Spring Boot", "Microservicios", "AWS", "Scrum"],
     "Ingeniera de sistemas con enfoque en arquitecturas de microservicios y nube."),
    ("Contador Publico", ["NIIF", "Siigo", "Excel avanzado", "Auditoria", "Nomina"],
     "Contador publico con experiencia en cierres contables y reportes tributarios."),
    ("Disenadora UX/UI", ["Figma", "Design Systems", "HTML", "CSS", "Investigacion de usuarios"],
     "Disenadora de experiencia de usuario centrada en productos digitales accesibles."),
    ("Administrador de Redes", ["Cisco", "Linux", "Firewalls", "VPN", "Monitoreo"],
     "Especialista en administracion de infraestructura de red y seguridad perimetral."),
    ("Asistente Administrativa", ["Gestion documental", "Atencion al cliente", "Office", "Archivo", "Agenda"],
     "Asistente administrativa con experiencia en gestion documental y servicio al cliente."),
    ("Ingeniero Electronico", ["PLC", "Automatizacion", "SCADA", "Mantenimiento", "IoT"],
     "Ingeniero electronico dedicado a la automatizacion de procesos industriales."),
    ("Especialista en Marketing Digital", ["SEO", "Google Ads", "Redes sociales", "Analitica web", "Copywriting"],
     "Especialista en marketing digital enfocado en crecimiento organico y pauta."),
    ("Tecnologo en Gestion Empresarial", ["Procesos", "Indicadores", "Excel", "Calidad", "Logistica"],
     "Tecnologo en gestion empresarial con enfoque en mejora continua de procesos."),
]

UNIVERSIDADES = [
    "Unidades Tecnologicas de Santander", "Universidad Industrial de Santander",
    "Universidad de Investigacion y Desarrollo", "Universidad Pontificia Bolivariana",
    "SENA - Servicio Nacional de Aprendizaje", "Universidad Cooperativa de Colombia",
]

IDIOMAS_POOL = [
    ["Espanol (nativo)", "Ingles (B1)"],
    ["Espanol (nativo)", "Ingles (B2)", "Portugues (A2)"],
    ["Espanol (nativo)", "Ingles (A2)"],
    ["Espanol (nativo)", "Ingles (C1)"],
    ["Espanol (nativo)"],
]


def build_resume(i: int) -> dict:
    nombre = f"{fake.first_name()} {fake.last_name()} {fake.last_name()}"
    cargo, habilidades, perfil_base = PERFILES[i]
    perfil = (
        f"{perfil_base} Cuenta con {rng.randint(2, 12)} anos de experiencia en el sector y "
        "capacidad de trabajo en equipo, comunicacion asertiva y aprendizaje continuo."
    )
    educacion = [
        f"{rng.choice(['Tecnologia', 'Pregrado', 'Tecnica profesional'])} en "
        f"{rng.choice(['Desarrollo de Software', 'Sistemas', 'Administracion', 'Electronica', 'Mercadeo'])} - "
        f"{rng.choice(UNIVERSIDADES)} ({rng.randint(2012, 2022)})",
    ]
    if i % 2 == 0:
        educacion.append(
            f"Diplomado en {rng.choice(['Analitica de datos', 'Seguridad informatica', 'Gestion de proyectos', 'Marketing digital'])} - "
            f"{rng.choice(UNIVERSIDADES)} ({rng.randint(2019, 2024)})"
        )
    experiencia = []
    for _ in range(rng.randint(1, 3)):
        empresa = fake.company()
        y1 = rng.randint(2014, 2021)
        y2 = y1 + rng.randint(1, 3)
        experiencia.append(
            f"{rng.choice([cargo, 'Auxiliar', 'Analista', 'Lider tecnico'])} en {empresa} ({y1} - {y2}): "
            f"{rng.choice(['apoyo a procesos internos', 'desarrollo de proyectos del area', 'atencion de requerimientos y mejoras', 'gestion de clientes y proveedores'])}."
        )
    idiomas = rng.choice(IDIOMAS_POOL)
    lines = [
        nombre,
        f"HOJA DE VIDA - {cargo}",
        f"Correo: {nombre.split()[0].lower()}.{nombre.split()[1].lower()}@correo-ficticio.example.com",
        f"Ciudad: {fake.city()}",
        "",
        "PERFIL PROFESIONAL",
        perfil,
        "",
        "EDUCACION",
    ]
    lines += [f"- {e}" for e in educacion]
    lines += ["", "EXPERIENCIA LABORAL"]
    lines += [f"- {e}" for e in experiencia]
    lines += ["", "HABILIDADES"]
    lines += [f"- {h}" for h in habilidades]
    lines += ["", "IDIOMAS"]
    lines += [f"- {idm}" for idm in idiomas]
    lines += ["", "Referencias disponibles a solicitud.",
              "Documento generado con datos ficticios para fines academicos."]
    return {
        "text": "\n".join(lines),
        "title": f"Hoja de vida - {nombre}",
        "ground_truth": {
            "category": "talento_humano",
            "fields": {
                "nombre": nombre,
                "num_educacion": len(educacion),
                "num_experiencia": len(experiencia),
                "num_habilidades": len(habilidades),
                "idiomas": idiomas,
            },
        },
    }


# ============================================================ writers
def write_txt(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def write_docx(path: Path, text: str, title: str) -> None:
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


def write_pdf(path: Path, text: str, title: str) -> None:
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer, pagesize=letter,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=title,
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=14)
    flow = [Paragraph(title, styles["Heading1"]), Spacer(1, 8)]
    for line in text.splitlines():
        if line.strip():
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            flow.append(Paragraph(safe, body))
        else:
            flow.append(Spacer(1, 8))
    pdf.build(flow)
    path.write_bytes(buffer.getvalue())


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    builders = [("factura", build_invoice), ("contrato", build_contract), ("hoja_de_vida", build_resume)]
    # Per category: 4 PDF, 3 DOCX, 3 TXT -> totals: 12 PDF, 9 DOCX, 9 TXT
    format_cycle = ["pdf", "pdf", "pdf", "pdf", "docx", "docx", "docx", "txt", "txt", "txt"]

    ground_truth = []
    for prefix, builder in builders:
        for i in range(10):
            data = builder(i)
            fmt = format_cycle[i]
            filename = f"{prefix}_{i + 1:02d}.{fmt}"
            path = OUTPUT_DIR / filename
            if fmt == "txt":
                write_txt(path, data["text"])
            elif fmt == "docx":
                write_docx(path, data["text"], data["title"])
            else:
                write_pdf(path, data["text"], data["title"])
            entry = {"filename": filename, "format": fmt, "title": data["title"]}
            entry.update(data["ground_truth"])
            ground_truth.append(entry)
            print(f"[corpus] {filename} generado")

    (OUTPUT_DIR / "ground_truth.json").write_text(
        json.dumps(ground_truth, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # RAG evaluation set: questions answerable from the corpus + one unanswerable
    invoices = [g for g in ground_truth if g["category"] == "financiero"]
    contracts = [g for g in ground_truth if g["category"] == "legal"]
    resumes = [g for g in ground_truth if g["category"] == "talento_humano"]
    inv0 = invoices[0]
    con0 = contracts[0]
    res0 = resumes[0]
    rag_eval = [
        {
            "id": "RAG-01",
            "question": f"¿Cual es el total a pagar de la factura {inv0['fields']['numero_factura']}?",
            "expected_contains": str(int(inv0["fields"]["total"])) if inv0["fields"]["moneda"] == "COP" else str(inv0["fields"]["total"]),
            "expected_source": inv0["filename"],
            "expect_grounded": True,
        },
        {
            "id": "RAG-02",
            "question": f"¿Cual es la fecha de terminacion del contrato entre {con0['fields']['partes'][0]} y {con0['fields']['partes'][1]}?",
            "expected_contains": con0["fields"]["fecha_terminacion"],
            "expected_source": con0["filename"],
            "expect_grounded": True,
        },
        {
            "id": "RAG-03",
            "question": f"¿Que habilidades tiene {res0['fields']['nombre']}?",
            "expected_contains": None,
            "expected_source": res0["filename"],
            "expect_grounded": True,
        },
        {
            "id": "RAG-04",
            "question": "¿Cual es la receta de la bandeja paisa tradicional?",
            "expected_contains": "No se encontro evidencia suficiente",
            "expected_source": None,
            "expect_grounded": False,
        },
    ]
    (OUTPUT_DIR / "rag_eval.json").write_text(
        json.dumps(rag_eval, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[corpus] {len(ground_truth)} documentos + ground_truth.json + rag_eval.json en {OUTPUT_DIR}")


if __name__ == "__main__":
    sys.exit(main())
